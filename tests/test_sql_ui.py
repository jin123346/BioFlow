from pathlib import Path
import os
import sys
import tempfile
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication, QDialog

from config.sql_media_profiles import default_profile
from services.sql_generation_service import auto_map_profile
from ui.tabs.sql_tab import SqlTab
from ui.widgets.sql_mapping_dialog import SQLMappingDialog


class SQLUITests(unittest.TestCase):
    def test_mapping_preview_resolves_files_before_schema_confirmation(self):
        image = self.root / "DB001.JPG"
        image.write_bytes(b"preview fixture")
        row = {"표본정보순번": "SMPL-001", "데이터베이스번호": "DB001", "기관코드": "00123"}
        profile = auto_map_profile(default_profile("specimen_image"), list(row))
        profile.update(media_lookup_dir=str(self.root))
        next(col for col in profile["columns"] if col["name"] == "STORE_FILE_NM")["folder_lookup"] = True
        dialog = SQLMappingDialog(profile, list(row))
        try:
            dialog.set_preview_rows([(4, row)])
            dialog.refresh_mapping_preview()
            table = dialog.mapping_preview
            self.assertEqual(table.item(0, 4).text(), "DB001.JPG")
            self.assertEqual(table.item(0, 5).text(), str(image))
            self.assertEqual(table.item(0, 7).text(), "/NAS/naris/smpl/naris/00123/DB001.JPG")
            self.assertEqual(table.item(0, 8).text(), "매핑 확인 (DB 검증 전)")
            self.assertTrue(dialog.values()["schema_confirmed"])
            (self.root / "DB001.png").write_bytes(b"another image")
            dialog.refresh_mapping_preview()
            self.assertIn("후보가 2개", table.item(0, 8).text())
            self.assertEqual(table.item(0, 4).text(), "")
            self.assertEqual(table.item(0, 7).text(), "")
        finally:
            dialog.close()

    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.settings_patch = patch("ui.tabs.sql_tab.SETTINGS_FILE", self.root / "settings.json")
        self.settings_patch.start()
        self.addCleanup(self.settings_patch.stop)
        self.errors = []
        self.hook = patch.object(sys, "excepthook", lambda kind, value, tb: self.errors.append(str(value)))
        self.hook.start()
        self.addCleanup(self.hook.stop)
        workbook = Workbook()
        workbook.active.title = "Media"
        workbook.active.append(["설명 제목"])
        workbook.active.append(["표본정보순번", "미디어타입", "파일명"])
        for row in range(30):
            workbook.active.append([row + 1, "이미지", f"image{row}.jpg"])
        other = workbook.create_sheet("Hidden")
        other["A1"] = "다른 시트"
        other.sheet_state = "hidden"
        self.source = self.root / "input.xlsx"
        workbook.save(self.source)
        workbook.close()

    def test_load_preview_generate_save_and_invalidate_on_sheet_change(self):
        tab = SqlTab()
        try:
            tab.load_file(self.source)
            self.assertEqual(tab.header_row.value(), 2)
            tab.header_row.setValue(2)
            self.assertEqual(tab.preview.rowCount(), 20)
            self.assertIn("30", tab.info.text())
            self.assertIn("숨김", tab.sheets.itemText(1))
            profile = auto_map_profile(default_profile("specimen_image"), tab.sheet_data.headers)
            profile.update(schema_confirmed=True, nas_template="/NAS/images/")
            for column in profile["columns"]:
                if column["type"] == "VARCHAR2":
                    column["length"] = 1000
            with patch("ui.tabs.sql_tab.QInputDialog.getItem", return_value=("표본이미지", True)), \
                 patch("ui.tabs.sql_tab.SQLMappingDialog") as dialog:
                dialog.return_value.exec.return_value = QDialog.Accepted
                dialog.return_value.values.return_value = profile
                tab.generate_button.click()
            self.assertEqual(len(tab.results), 30)
            self.assertTrue(all(item.status == "SUCCESS" for item in tab.results))
            self.assertIn("INSERT INTO", tab.sql_preview.toPlainText())
            self.assertTrue(tab.save_button.isEnabled())
            self.assertTrue((self.root / "settings.json").exists())
            with patch("ui.tabs.sql_tab.QFileDialog.getSaveFileName", return_value=(str(self.root / "out.sql"), "")), \
                 patch("ui.tabs.sql_tab.QMessageBox.information") as info, \
                 patch("ui.tabs.sql_tab.QMessageBox.critical") as error:
                tab.save_button.click()
                error.assert_not_called()
                info.assert_called_once()
            self.assertTrue((self.root / "out.sql").is_file())
            self.assertTrue((self.root / "out_검수.xlsx").is_file())
            tab.sheets.setCurrentIndex(1)
            self.assertIsNone(tab.results)
            self.assertFalse(tab.save_button.isEnabled())
            self.assertEqual(self.errors, [])
        finally:
            tab.close()

    def test_mapping_dialog_exposes_required_thumbnail_and_unknown_schema(self):
        for kind in ("observation_image", "specimen_image", "observation_video", "observation_sound"):
            profile = auto_map_profile(default_profile(kind), ["종고유ID", "파일명", "썸네일 파일명"])
            dialog = SQLMappingDialog(profile, ["종고유ID", "파일명", "썸네일 파일명"])
            try:
                result = dialog.values()
                self.assertTrue(result["schema_confirmed"])
                self.assertEqual(result["table"], profile["table"])
                if "image" in kind:
                    self.assertEqual(result["sequence_value"], "1")
                    seq_index = next(i for i, col in enumerate(result["columns"]) if col["name"] == "IMAGE_SEQ")
                    seq = result["columns"][seq_index]
                    self.assertEqual((seq["type"], seq["length"], seq["constant"]), ("CHAR", 1, "1"))
                    self.assertTrue(all(not control.isEnabled() for control in dialog.controls[seq_index]))
                    self.assertEqual(result["value_codes"], ["03"])
                if kind == "observation_video":
                    thumb = next(col for col in result["columns"] if col["name"] == "LVBNG_MVPCT_RPRSN_IMAGE")
                    self.assertEqual(thumb["length"], 100)
                    self.assertEqual(thumb["source"], "썸네일 파일명")
                    self.assertTrue(thumb["required"])
            finally:
                dialog.close()
        self.assertEqual(self.errors, [])

    def test_sql_page_is_lazy_loaded_and_does_not_open_upload_page(self):
        from ui.main_window import MainWindow
        window = MainWindow()
        try:
            self.assertEqual(window.nav_list.item(4).text(), "SQL 생성")
            window.nav_list.setCurrentRow(4)
            self.assertIsInstance(window.pages.currentWidget(), SqlTab)
            self.assertNotIn(2, window._loaded_pages)
            self.assertEqual(self.errors, [])
        finally:
            window.close()

    def test_filename_extension_is_editable_and_saved_per_mapping_column(self):
        headers = ["데이터베이스번호", "종고유ID"]
        profile = auto_map_profile(default_profile("observation_video"), headers)
        dialog = SQLMappingDialog(profile, headers)
        try:
            for column, controls in zip(profile["columns"], dialog.controls):
                if column["name"] in ("PH500_FILE_NM", "LVBNG_MVPCT_RPRSN_IMAGE"):
                    controls[0].setCurrentIndex(controls[0].findData("데이터베이스번호"))
            dialog.extension_controls["PH500_FILE_NM"].setCurrentText(".mp4")
            dialog.extension_controls["LVBNG_MVPCT_RPRSN_IMAGE"].setCurrentText(".jpg")
            result = dialog.values()
            movie = next(col for col in result["columns"] if col["name"] == "PH500_FILE_NM")
            thumb = next(col for col in result["columns"] if col["name"] == "LVBNG_MVPCT_RPRSN_IMAGE")
            self.assertEqual(movie["source"], "데이터베이스번호")
            self.assertEqual(movie["filename_extension"], ".mp4")
            self.assertEqual(thumb["filename_extension"], ".jpg")
            self.assertFalse(dialog.extension_controls["SPECS_ESSNT_ID"].isEnabled())
            restored = SQLMappingDialog(result, headers)
            try:
                self.assertEqual(restored.extension_controls["PH500_FILE_NM"].currentText(), ".mp4")
            finally:
                restored.close()
        finally:
            dialog.close()

    def test_folder_matching_options_are_saved_and_disable_excel_filename_source(self):
        headers = ["데이터베이스번호", "파일명"]
        profile = auto_map_profile(default_profile("specimen_image"), headers)
        dialog = SQLMappingDialog(profile, headers)
        try:
            dialog.lookup_folder.setText(str(self.root))
            dialog.lookup_controls["STORE_FILE_NM"].setCurrentIndex(1)
            index = next(index for index, col in enumerate(profile["columns"]) if col["name"] == "STORE_FILE_NM")
            self.assertFalse(dialog.controls[index][0].isEnabled())
            result = dialog.values()
            self.assertEqual(result["media_lookup_dir"], str(self.root))
            self.assertEqual(result["database_number_column"], "데이터베이스번호")
            self.assertTrue(next(col for col in result["columns"] if col["name"] == "STORE_FILE_NM")["folder_lookup"])
            self.assertTrue(result["lookup_recursive"])
            self.assertTrue(result["lookup_skip_previews"])
            restored = SQLMappingDialog(result, headers)
            try:
                self.assertTrue(restored.lookup_controls["STORE_FILE_NM"].currentData())
                self.assertEqual(restored.lookup_folder.text(), str(self.root))
            finally:
                restored.close()
        finally:
            dialog.close()

    def test_completed_sheet_defaults_to_selected_type_when_filter_columns_are_absent(self):
        headers = ["종고유ID", "이전데이터베이스번호", "학명"]
        profile = auto_map_profile(default_profile("specimen_image"), headers)
        dialog = SQLMappingDialog(profile, headers)
        try:
            self.assertTrue(dialog.values()["all_rows_confirmed"])
            self.assertEqual(dialog.values()["database_number_column"], "이전데이터베이스번호")
            self.assertIn("전체 행", dialog.filter_notice.text())
            self.assertIn("표본이미지", dialog.filter_notice.text())
            dialog.type_column.addItem("미디어타입", "미디어타입")
            dialog.type_column.setCurrentIndex(dialog.type_column.count() - 1)
            self.assertFalse(dialog.values()["all_rows_confirmed"])
        finally:
            dialog.close()


if __name__ == "__main__":
    unittest.main()
