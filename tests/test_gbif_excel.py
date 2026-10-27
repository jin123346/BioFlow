import tempfile
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from services.gbif_excel_service import GBIF_HEADERS, country_code, write_gbif_template
from ui.tabs.upload_tab import UploadTab


class GbifExcelTests(unittest.TestCase):
    def record(self, specimen=False):
        location = "채집위치" if specimen else "관찰위치"
        return {
            "데이터베이스번호": "TEST-SP-BI-260300100",
            "기관코드": "TEST", "학명": "Sus scrofa Linnaeus, 1758",
            "국명": "멧돼지", "분류코드": "BI", "위도": "36.58675", "경도": "127.101611",
            "국가국문명": "대한민국", "국가영문명": "Republic of Korea",
            "동정년월일": "261015", "동정자 영문": "Test Identifier",
            "채집년월일" if specimen else "관찰일": "20261014",
            f"{location} 시도 영문": "Chungcheongnam-do",
            f"{location} 시군구 영문": "Gongju-si",
            f"{location} 읍면동 영문": "Jeongan-myeon",
            f"{location} 상세 영문": "Goseong-ri",
        }

    def test_both_kinds_match_supplied_layout(self):
        for specimen in (False, True):
            with self.subTest(specimen=specimen), tempfile.TemporaryDirectory() as folder:
                path = Path(folder) / "export.xlsx"
                kind = "specimen" if specimen else "observation"
                record = self.record(specimen)
                write_gbif_template(path, [record], kind)
                workbook = load_workbook(path)
                try:
                    sheet = workbook.worksheets[0]
                    self.assertEqual(tuple(c.value for c in sheet[1]), GBIF_HEADERS)
                    self.assertEqual(sheet.max_row, 2)
                    values = {name: sheet.cell(2, col).value for col, name in enumerate(GBIF_HEADERS, 1)}
                    self.assertEqual(values["occurrenceID"], record["데이터베이스번호"])
                    self.assertEqual(values["catalogNumber"], record["데이터베이스번호"])
                    self.assertEqual(values["basisOfRecord"], "PreservedSpecimen" if specimen else "HumanObservation")
                    self.assertEqual(values["eventDate"].date(), date(2026, 10, 14))
                    self.assertEqual(values["Identification Date"].date(), date(2026, 10, 15))
                    self.assertEqual(values["scientificName"], record["학명"])
                    self.assertEqual(values["collectionCode"], "BI")
                    self.assertEqual(values["decimalLatitude"], 36.58675)
                    self.assertEqual(values["decimalLongitude"], 127.101611)
                    self.assertEqual(values["countryCode"], "KR")
                    self.assertEqual(values["Identified by"], "Test Identifier")
                    self.assertEqual(values["stateProvince"], "Chungcheongnam-do")
                    self.assertEqual(values["county"], "Gongju-si")
                    self.assertEqual(values["locality"], "Jeongan-myeon")
                    self.assertEqual(values["verbatimLocality"], "Goseong-ri")
                    self.assertEqual(values["vernacularName"], "멧돼지")
                finally:
                    workbook.close()

    def test_no_example_leaks_missing_values_or_formulas(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "export.xlsx"
            write_gbif_template(path, [{"국명": "=1+1"}], "observation")
            workbook = load_workbook(path)
            try:
                sheet = workbook.worksheets[0]
                self.assertIsNone(sheet["A2"].value)
                self.assertIsNone(sheet["I2"].value)
                self.assertIsNone(sheet["K2"].value)
                self.assertEqual(sheet["Q2"].value, "=1+1")
                self.assertEqual(sheet["Q2"].data_type, "s")
            finally:
                workbook.close()

    def test_country_mapping_does_not_assume_korea(self):
        self.assertEqual(country_code({"국가영문명": "Japan"}), "JP")
        self.assertEqual(country_code({"국가국문명": "일본"}), "JP")
        self.assertEqual(country_code({"국가영문명": "US"}), "US")
        self.assertEqual(country_code({"국가영문명": "Unknown", "국가국문명": "대한민국"}), "")
        self.assertEqual(country_code({}), "")

    def test_large_export_extends_correct_coordinate_and_date_ranges(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "export.xlsx"
            write_gbif_template(path, [self.record()] * 1800, "observation")
            workbook = load_workbook(path)
            try:
                sheet = workbook.worksheets[0]
                self.assertEqual(sheet.max_row, 1801)
                self.assertEqual(sheet["Q1801"].value, "멧돼지")
                ranges = {str(v.sqref): (v.formula1, v.formula2) for v in sheet.data_validations.dataValidation}
                self.assertEqual(ranges["G2:G1801"], ("-90", "90"))
                self.assertEqual(ranges["H2:H1801"], ("-180", "180"))
                self.assertEqual(ranges["C2:C1801"], ("DATE(1900,1,1)", "DATE(9999,12,31)"))
            finally:
                workbook.close()

    def test_both_export_actions_create_three_files_and_keep_form_on_failure(self):
        for specimen in (False, True):
            for fail in (False, True):
                with self.subTest(specimen=specimen, fail=fail), tempfile.TemporaryDirectory() as folder:
                    path = Path(folder) / "records.xlsx"
                    ui = SimpleNamespace()
                    ui.final_excel_path = UploadTab.final_excel_path
                    ui.gbif_excel_path = UploadTab.gbif_excel_path
                    ui.table_to_records = Mock(return_value=[self.record(specimen)])
                    ui.confirm_location_issues_before_save = Mock(return_value=True)
                    ui.prepare_media_for_excel_export = Mock(return_value=True)
                    for name in (
                        "sync_current_specimen_form_to_table", "copy_pending_media_for_all_specimen_rows",
                        "sync_current_form_to_table", "copy_pending_media_for_all_rows", "capture_pending_video_thumbnails",
                        "reset_after_specimen_export", "reset_after_export",
                    ):
                        setattr(ui, name, Mock())
                    ui.data_table = ui.specimen_data_table = object()
                    ui.write_excel_file = UploadTab.write_excel_file
                    ui.write_final_excel_file = UploadTab.write_final_excel_file
                    ui.write_specimen_excel_file = UploadTab.write_specimen_excel_file
                    ui.write_specimen_final_excel_file = UploadTab.write_specimen_final_excel_file
                    action = UploadTab.export_specimen_rows_to_excel if specimen else UploadTab.export_rows_to_excel
                    with patch("ui.tabs.upload_tab.QFileDialog.getSaveFileName", return_value=(str(path), "")), \
                         patch("ui.tabs.upload_tab.QMessageBox.information") as info, \
                         patch("ui.tabs.upload_tab.QMessageBox.critical") as error, \
                         patch("ui.tabs.upload_tab.write_gbif_template", side_effect=OSError("test failure") if fail else write_gbif_template):
                        action(ui, include_media=True)
                    ui.prepare_media_for_excel_export.assert_called_once_with(specimen, choose_destination=True)
                    self.assertFalse(ui._defer_media_copy)
                    reset = ui.reset_after_specimen_export if specimen else ui.reset_after_export
                    if fail:
                        reset.assert_not_called()
                        error.assert_called_once()
                        info.assert_not_called()
                    else:
                        reset.assert_called_once()
                        info.assert_called_once()
                        error.assert_not_called()
                        self.assertEqual(len(list(Path(folder).glob("*.xlsx"))), 3)
                        self.assertTrue(ui.gbif_excel_path(path).exists())

    def test_excel_only_export_skips_media_and_keeps_editor_rows(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "records.xlsx"
            ui = SimpleNamespace(
                data_table=object(),
                _defer_media_copy=False,
                sync_current_form_to_table=Mock(),
                table_to_records=Mock(return_value=[self.record(False)]),
                confirm_location_issues_before_save=Mock(return_value=True),
                prepare_media_for_excel_export=Mock(return_value=True),
                reset_after_export=Mock(),
                final_excel_path=UploadTab.final_excel_path,
                gbif_excel_path=UploadTab.gbif_excel_path,
                write_excel_file=UploadTab.write_excel_file,
                write_final_excel_file=UploadTab.write_final_excel_file,
            )
            with patch("ui.tabs.upload_tab.QFileDialog.getSaveFileName", return_value=(str(path), "")), \
                 patch("ui.tabs.upload_tab.QMessageBox.information"), \
                 patch("ui.tabs.upload_tab.QMessageBox.critical"):
                UploadTab.export_rows_to_excel(ui, include_media=False)
            ui.prepare_media_for_excel_export.assert_not_called()
            ui.reset_after_export.assert_not_called()
            self.assertTrue(path.exists())
            self.assertTrue(ui.final_excel_path(path).exists())
            self.assertTrue(ui.gbif_excel_path(path).exists())

    def test_manual_working_save_writes_one_basic_excel_without_media(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "working.xlsx"
            ui = SimpleNamespace(
                data_table=object(),
                sync_current_form_to_table=Mock(),
                table_to_records=Mock(return_value=[self.record(False)]),
                write_excel_file=UploadTab.write_excel_file,
                write_draft=Mock(),
                collect_form_values=Mock(return_value={}),
                set_save_status=Mock(),
            )
            with patch("ui.tabs.upload_tab.QFileDialog.getSaveFileName", return_value=(str(path), "")), \
                 patch("ui.tabs.upload_tab.QMessageBox.information"):
                UploadTab.save_draft(ui)
            self.assertTrue(path.exists())
            self.assertEqual(len(list(Path(folder).glob("*.xlsx"))), 1)
            ui.write_draft.assert_called_once()


if __name__ == "__main__":
    unittest.main()
