import json
import os
from contextlib import ExitStack
from pathlib import Path
from unittest.mock import Mock, patch

import pytest

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
from PySide6.QtCore import QTimer
from PySide6.QtWidgets import QApplication, QDialog

from services.address_mapping_service import METADATA_KEY, apply_address, map_address, valid_metadata
from services.vworld_geocoding_service import GeocodingError, VWorldAddress
from ui.dialogs.address_lookup_dialog import AddressLookupDialog, AddressPreviewDialog, AddressScopeDialog
from ui.tabs.upload_tab import ADDRESS_FIELDS, SPECIMEN_ADDRESS_FIELDS, ROW_COLUMNS, SPECIMEN_ROW_COLUMNS, UploadTab

CASES = json.loads((Path(__file__).parent / "fixtures/vworld_live_addresses.json").read_text(encoding="utf-8"))["cases"]


def address(name="anyang"):
    case = next(row for row in CASES if row["name"] == name)
    return VWorldAddress(case["structure"], case["text"], case["latitude"], case["longitude"])


@pytest.mark.parametrize("name,city,district,dong", [
    ("seoul", "", "중구", "태평로1가"), ("anyang", "안양시", "동안구", "관양동"),
    ("dongnae", "", "동래구", "복천동"), ("gyeongju", "", "경주시", "동천동"),
    ("namhae", "", "남해군", "남해읍"),
])
def test_captured_live_mapping(name, city, district, dong):
    fields, metadata = map_address(address(name))
    assert metadata["city"] == city
    assert metadata["district"] == district
    assert fields["관찰위치 시군구 국문"] == " ".join(filter(None, (city, district)))
    assert fields["관찰위치 읍면동 국문"] == dong


def test_separate_levels_and_detail_fallback():
    result = VWorldAddress({"level2": "안양시", "level3": "동안구", "detail": "보조", "level4LC": "00123"}, "", "37", "127")
    fields, metadata = map_address(result)
    assert fields["관찰위치 시군구 국문"] == "안양시 동안구"
    assert fields["관찰위치 상세 국문"] == "보조"
    assert metadata["legal_dong_code"] == "00123"
    result.structure["level5"] = "12-1"
    assert map_address(result)[0]["관찰위치 상세 국문"] == "12-1"


def test_overwrite_missing_fields_and_stale_metadata():
    result = address()
    original = {"위도": result.latitude, "경도": result.longitude,
                "관찰위치 상세 국문": "현장 기록", "관찰위치 상세 영문": "Keep"}
    updated = apply_address(original, result)
    assert updated["관찰위치 상세 국문"] == "현장 기록"
    assert not updated[METADATA_KEY]["matches_displayed_address"]
    assert original.get("관찰위치 시도 국문") is None
    replaced = apply_address(original, result, overwrite=True)
    assert replaced["관찰위치 상세 국문"] == "1590"
    assert replaced["관찰위치 상세 영문"] == "Keep"
    assert valid_metadata(replaced, replaced[METADATA_KEY])
    replaced["위도"] = "38"
    assert valid_metadata(replaced, replaced[METADATA_KEY]) is None
    partial = VWorldAddress({"level1": "경기도"}, "", result.latitude, result.longitude)
    assert apply_address(original, partial, overwrite=True)["관찰위치 상세 국문"] == "현장 기록"


@pytest.fixture
def app():
    return QApplication.instance() or QApplication([])


@pytest.fixture
def editor(app):
    with ExitStack() as stack:
        for method in ("_ensure_reference_files", "load_reference_data", "load_species_master_lookup",
                       "load_draft", "load_specimen_draft", "schedule_autosave", "schedule_specimen_autosave",
                       "autosave_draft", "autosave_specimen_draft"):
            stack.enter_context(patch.object(UploadTab, method))
        widget = UploadTab()
        yield widget
        widget.close()


@pytest.mark.parametrize("specimen", [False, True])
def test_buttons_apply_korean_preserve_english_and_metadata_roundtrip(editor, specimen):
    result = address()
    prefix = "채집위치" if specimen else "관찰위치"
    table = editor.specimen_data_table if specimen else editor.data_table
    columns = SPECIMEN_ROW_COLUMNS if specimen else ROW_COLUMNS
    widgets = editor.specimen_form_widgets if specimen else editor.form_widgets
    add = editor.add_specimen_data_row if specimen else editor.add_data_row
    collect = editor.collect_specimen_form_values if specimen else editor.collect_form_values
    add({"NO.": 1, "국명": "테스트", "위도": result.latitude, "경도": result.longitude,
         f"{prefix} 상세 국문": "현장 기록", f"{prefix} 시도 영문": "Keep"})
    table.selectRow(0)

    def lookup(parent, targets, is_specimen):
        assert is_specimen == specimen
        return {i: apply_address(record, result, specimen) for i, record in targets}

    with patch("ui.tabs.upload_tab.lookup_addresses", side_effect=lookup):
        (editor.fill_specimen_address_from_latlon if specimen else editor.fill_address_from_latlon)()
    assert widgets[f"{prefix} 시군구 국문"].text() == "안양시 동안구"
    assert widgets[f"{prefix} 상세 국문"].text() == "현장 기록"
    assert widgets[f"{prefix} 시도 영문"].text() == "Keep"
    record = editor.table_to_records(table, columns)[0]
    assert record[METADATA_KEY]["legal_dong_code"] == "4117310200"
    assert collect()[METADATA_KEY] == record[METADATA_KEY]
    # JSON serialization (drafts), reordering and table reconstruction retain metadata.
    restored = json.loads(json.dumps(editor.editor_table_records(specimen)))
    editor.replace_editor_table_records(specimen, restored)
    assert editor.table_to_records(table, columns)[0][METADATA_KEY] == record[METADATA_KEY]
    widgets["위도"].setText("38")
    assert METADATA_KEY not in collect()
    (editor.clear_specimen_form if specimen else editor.clear_form)()
    assert widgets["위도"].property(METADATA_KEY) is None


@pytest.mark.parametrize("specimen", [False, True])
def test_batch_partial_failure_and_selected_form_sync(editor, specimen):
    result = address()
    table = editor.specimen_data_table if specimen else editor.data_table
    columns = SPECIMEN_ROW_COLUMNS if specimen else ROW_COLUMNS
    add = editor.add_specimen_data_row if specimen else editor.add_data_row
    fields = SPECIMEN_ADDRESS_FIELDS if specimen else ADDRESS_FIELDS
    for i in range(2):
        add({"NO.": i + 1, "위도": result.latitude, "경도": result.longitude, "국명": "테스트"})
    table.selectRow(0)
    before = editor.address_row_record(table, columns, 1)
    def lookup(parent, targets, is_specimen, differences_only=False):
        return {0: apply_address(targets[0][1], result, specimen)}
    with patch("ui.tabs.upload_tab.lookup_addresses", side_effect=lookup):
        editor.batch_fill_addresses(table, columns, fields, Mock())
    assert editor.address_row_record(table, columns, 1) == before
    prefix = "채집위치" if specimen else "관찰위치"
    widgets = editor.specimen_form_widgets if specimen else editor.form_widgets
    assert widgets[f"{prefix} 시군구 국문"].text() == "안양시 동안구"
    assert METADATA_KEY in editor.table_to_records(table, columns)[0]


def test_stale_input_not_applied(editor):
    result = address()
    editor.set_form_value("위도", result.latitude)
    editor.set_form_value("경도", result.longitude)
    def lookup(parent, targets, specimen):
        editor.set_form_value("위도", "38")
        return {0: apply_address(targets[0][1], result)}
    with patch("ui.tabs.upload_tab.lookup_addresses", side_effect=lookup):
        editor.fill_address_from_latlon()
    assert editor.get_form_value("관찰위치 시도 국문") == ""


def test_worker_failure_stops_remaining_requests(app):
    service = Mock()
    service.lookup.side_effect = GeocodingError("authentication", "인증 실패")
    dialog = AddressLookupDialog(service, "synthetic", [(0, {}), (1, {})])
    assert dialog.exec() == QDialog.Accepted
    assert service.lookup.call_count == 1
    assert dialog.worker.results == {}
    assert dialog.worker.key == ""


def test_cancel_discard_results_and_finish_thread(app):
    import threading
    release = threading.Event()
    service = Mock()
    service.lookup.side_effect = lambda *args: (release.wait(2), address())[1]
    dialog = AddressLookupDialog(service, "synthetic", [(0, {}), (1, {})])
    def cancel():
        dialog.reject()
        release.set()
    QTimer.singleShot(50, cancel)
    assert dialog.exec() == QDialog.Rejected
    assert dialog.worker.isFinished()
    assert dialog.worker.results == {}


def test_preview_defaults_to_preserving_existing_values(app):
    result = address()
    record = {"관찰위치 상세 국문": "현장 기록"}
    dialog = AddressPreviewDialog([(0, record)], {0: result}, {})
    assert not dialog.overwrite.isChecked()
    assert dialog.table.item(4, 4).text() == "현장 기록"
    dialog.overwrite.setChecked(True)
    assert dialog.table.item(4, 4).text() == "1590"
    dialog.close()


def test_address_scope_defaults_to_empty_and_enables_selected_rows(app):
    empty = AddressScopeDialog(0)
    assert empty.mode() == "empty"
    assert not empty.selected.isEnabled()
    empty.close()
    selected = AddressScopeDialog(2)
    selected.selected.setChecked(True)
    assert selected.mode() == "selected"
    selected.all_rows.setChecked(True)
    assert selected.mode() == "all"
    selected.close()


def test_conflict_preview_shows_only_changed_existing_values_and_allows_edit(app):
    result = address()
    field = "관찰위치 상세 국문"
    record = {field: "현장 기록"}
    dialog = AddressPreviewDialog([(0, record)], {0: result}, {}, differences_only=True)
    assert dialog.table.rowCount() == 1
    assert dialog.table.item(0, 1).text() == field
    dialog.table.item(0, 4).setText("사용자 확인값")
    planned = dialog.planned_records()[0]
    assert planned[field] == "사용자 확인값"
    assert field in planned[METADATA_KEY]["write_fields"]
    dialog.close()


@pytest.mark.parametrize("specimen", [False, True])
def test_address_draft_and_excel_files_preserve_public_schema(editor, tmp_path, specimen):
    from openpyxl import load_workbook
    from services.excel_service import write_records
    result = address()
    record = apply_address({"NO.": 1, "위도": result.latitude, "경도": result.longitude}, result, specimen)
    columns = SPECIMEN_ROW_COLUMNS if specimen else ROW_COLUMNS
    draft_path = tmp_path / "draft.json"
    variable = "SPECIMEN_DRAFT_FILE" if specimen else "DRAFT_FILE"
    with patch(f"ui.tabs.upload_tab.{variable}", draft_path):
        writer = editor.write_specimen_draft if specimen else editor.write_draft
        reader = editor.read_specimen_draft if specimen else editor.read_draft
        writer([record], record)
        restored = reader()
    assert restored["rows"][0][METADATA_KEY]["legal_dong_code"] == "4117310200"
    assert restored["current_form"][METADATA_KEY] == record[METADATA_KEY]
    path = tmp_path / "export.xlsx"
    write_records(path, restored["rows"], columns)
    workbook = load_workbook(path)
    try:
        headers = [cell.value for cell in workbook.active[1]]
        assert headers == columns
        prefix = "채집위치" if specimen else "관찰위치"
        assert workbook.active.cell(2, headers.index(f"{prefix} 시군구 국문") + 1).value == "안양시 동안구"
        assert METADATA_KEY not in headers
    finally:
        workbook.close()


def test_first_use_opens_settings_then_reuses_saved_key(app):
    from controllers.address_controller import lookup_addresses
    key_service = Mock()
    key_service.load.side_effect = [("", None), ("synthetic", None)]
    result = address()
    with patch("controllers.address_controller.ApiKeyService", return_value=key_service), \
         patch("controllers.address_controller.ApiSettingsDialog") as settings, \
         patch("controllers.address_controller.AddressLookupDialog") as lookup, \
         patch("controllers.address_controller.AddressPreviewDialog") as preview:
        lookup.return_value.exec.return_value = QDialog.Accepted
        lookup.return_value.worker.results = {0: result}
        lookup.return_value.worker.errors = {}
        lookup.return_value.worker.references = {}
        lookup.return_value.worker.reference_errors = {}
        preview.return_value.exec.return_value = QDialog.Accepted
        preview.return_value.overwrite.isChecked.return_value = False
        preview.return_value.planned_records.return_value = {
            0: apply_address({"위도": result.latitude, "경도": result.longitude}, result)
        }
        updated = lookup_addresses(None, [(0, {"위도": result.latitude, "경도": result.longitude})])
        settings.return_value.exec.assert_called_once()
        assert lookup.call_args.args[1] == "synthetic"
        assert updated[0]["관찰위치 시군구 국문"] == "안양시 동안구"
