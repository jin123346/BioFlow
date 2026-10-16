from copy import copy
from pathlib import Path
import os
import struct
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from config.sql_media_profiles import default_profile
from services.sql_excel_service import ExcelSource, export_sql_review
from services.sql_generation_service import SQLResult, auto_map_profile, generate_sql


class SQLExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Media"
        sheet.append(["원본 제목"])
        sheet.append(["표본정보순번", "미디어타입", "파일명", "원본파일명", "기타 정보"])
        sheet.append([29848, "이미지", "image.jpg", "original", "보존할 값"])
        sheet.append([None, "이미지", "error.jpg", "original", "ID 누락"])
        sheet.append([12, "동영상", "video.mp4", "original", "제외 행"])
        sheet["A2"].fill = PatternFill("solid", fgColor="00AA00")
        other = workbook.create_sheet("Keep")
        other.append(["원본 수식", "=1+2", 42])
        other.sheet_state = "hidden"
        workbook.save(self.path)
        workbook.close()
        self.source = ExcelSource(self.path)
        self.sheet = self.source.read_sheet("Media", 2)
        self.profile = auto_map_profile(default_profile("specimen_image"), self.sheet.headers)
        self.profile.update(schema_confirmed=True, nas_template="/NAS/images/")
        for column in self.profile["columns"]:
            if column["type"] in ("VARCHAR2", "CHAR") and column["length"] is None:
                column["length"] = 1000
        self.results = generate_sql(self.sheet, self.profile)

    def test_review_preserves_all_original_cells_sheets_styles_and_source(self):
        before = self.path.read_bytes()
        txt, review = export_sql_review(self.source, self.sheet, self.results, self.profile, self.root / "out.txt")
        self.assertEqual(self.path.read_bytes(), before)
        text = txt.read_text(encoding="utf-8")
        self.assertEqual(text.count("INSERT INTO"), 1)
        self.assertIn(self.results[0].sql, text)
        original, output = load_workbook(self.path), load_workbook(review)
        try:
            self.assertEqual(output.sheetnames, original.sheetnames)
            for name in original.sheetnames:
                for row in original[name]:
                    for cell in row:
                        self.assertEqual(output[name][cell.coordinate].value, cell.value)
                        self.assertEqual(output[name][cell.coordinate]._style, cell._style)
            self.assertEqual(output["Keep"].sheet_state, "hidden")
            self.assertEqual(output["Keep"].max_column, 3)
            self.assertEqual(output["Media"]["F2"].value, "INSERT_SQL")
            self.assertEqual(output["Media"]["F3"].value, self.results[0].sql)
            self.assertNotIn("\n", output["Media"]["F3"].value)
            self.assertEqual(output["Media"].column_dimensions["F"].width, 40)
            self.assertEqual(output["Media"].column_dimensions["H"].width, 40)
            for coord in ("F2", "F3", "G3", "H4"):
                self.assertFalse(output["Media"][coord].alignment.wrap_text)
            for row_number in (2, 3, 4, 5):
                self.assertEqual(output["Media"].row_dimensions[row_number].height,
                                 original["Media"].row_dimensions[row_number].height)
            self.assertEqual(output["Media"]["G3"].value, "SUCCESS")
            self.assertEqual(output["Media"]["G4"].value, "ERROR")
            self.assertEqual(output["Media"]["G5"].value, "SKIP")
            self.assertIsNone(output["Media"]["F4"].value)
            self.assertIn("엑셀 4행", output["Media"]["H4"].value)
        finally:
            original.close()
            output.close()

    def test_sheet_snapshot_tracks_excel_rows_and_hidden_sheets(self):
        self.assertEqual(self.source.sheets, [("Media", False), ("Keep", True)])
        self.assertEqual([row for row, _ in self.sheet.rows], [3, 4, 5])
        self.assertEqual(self.source.suggest_header_row("Media"), 2)
        with self.assertRaises(ValueError):
            self.source.read_sheet("Media", 999)

    def test_real_completed_naris_templates_generate_image_sql_without_type_columns(self):
        from services.final_excel_service import write_final_template
        from ui.tabs.upload_tab import FINAL_EXPORT_COLUMNS, SPECIMEN_FINAL_EXPORT_COLUMNS
        for kind, columns in (("observation", FINAL_EXPORT_COLUMNS), ("specimen", SPECIMEN_FINAL_EXPORT_COLUMNS)):
            with self.subTest(kind=kind):
                path = self.root / f"{kind}_final.xlsx"
                database_number = "TEST-OB-BI-260100001"
                write_final_template(path, [{"종고유ID": "29848", "데이터베이스번호": database_number, "학명": "Test species"}], columns, kind)
                if kind in ("observation", "specimen"):
                    # This is an existing DB observation key, not a substitution with species ID.
                    workbook = load_workbook(path)
                    workbook.active.cell(3, 1, 777)
                    if kind == "specimen":
                        col = workbook.active.max_column + 1
                        workbook.active.cell(2, col, "SMPL_INFO_SEQ")
                        workbook.active.cell(3, col, "SMPL-001")
                    workbook.save(path)
                    workbook.close()
                source = ExcelSource(path)
                name = source.sheets[0][0]
                self.assertEqual(source.suggest_header_row(name), 2)
                sheet = source.read_sheet(name, 2)
                profile = auto_map_profile(default_profile(f"{kind}_image"), sheet.headers)
                self.assertTrue(profile["all_rows_confirmed"])
                folder = self.root / "media"
                folder.mkdir(exist_ok=True)
                (folder / (database_number + ".jpg")).write_bytes(b"image")
                profile.update(schema_confirmed=True, nas_template="/NAS/images/", media_lookup_dir=str(folder))
                for column in profile["columns"]:
                    if column["type"] == "VARCHAR2":
                        column["length"] = 1000
                    if column["name"] == "STORE_FILE_NM":
                        column["folder_lookup"] = True
                results = generate_sql(sheet, profile)
                self.assertEqual(len(results), 1)
                self.assertEqual(results[0].status, "SUCCESS", results[0].error)
                self.assertEqual(results[0].excel_row, 3)
                self.assertIn(database_number + ".jpg", results[0].sql)

    def test_uncalculated_mapped_formula_is_error_instead_of_silent_null(self):
        workbook = load_workbook(self.path)
        workbook["Media"]["D3"] = '=CONCAT("original", "name")'
        workbook.save(self.path)
        workbook.close()
        source = ExcelSource(self.path)
        sheet = source.read_sheet("Media", 2)
        results = generate_sql(sheet, self.profile)
        self.assertEqual(results[0].status, "ERROR")
        self.assertIn("수식 계산값", results[0].error)

    def test_original_and_existing_outputs_are_protected(self):
        with self.assertRaises(ValueError):
            export_sql_review(self.source, self.sheet, self.results, self.profile, self.path)
        # Even a separate TXT filename must not resolve its review output to the source.
        protected = self.root / "out_검수.xlsx"
        protected.write_bytes(self.path.read_bytes())
        source = ExcelSource(protected)
        with self.assertRaises(ValueError):
            export_sql_review(source, self.sheet, self.results, self.profile, self.root / "out.txt", overwrite=True)
        destination = self.root / "exists.txt"
        destination.write_text("old", encoding="utf-8")
        with self.assertRaises(FileExistsError):
            export_sql_review(self.source, self.sheet, self.results, self.profile, destination)
        self.assertEqual(destination.read_text(), "old")

    def test_second_output_failure_rolls_back_first_output(self):
        text = self.root / "out.txt"
        review = self.root / "out_검수.xlsx"
        text.write_bytes(b"previous text")
        review.write_bytes(b"previous workbook")
        replace = os.replace
        def fail_review(source, target):
            if Path(target) == review:
                raise PermissionError("review locked")
            return replace(source, target)
        with patch("services.sql_excel_service.os.replace", side_effect=fail_review):
            with self.assertRaises(PermissionError):
                export_sql_review(self.source, self.sheet, self.results, self.profile, text, overwrite=True)
        self.assertEqual(text.read_bytes(), b"previous text")
        self.assertEqual(review.read_bytes(), b"previous workbook")

    def test_oversized_sql_is_not_silently_truncated_by_excel(self):
        results = [SQLResult(3, "SUCCESS", "X" * 32768), *self.results[1:]]
        with self.assertRaisesRegex(ValueError, "32767"):
            export_sql_review(self.source, self.sheet, results, self.profile, self.root / "out.txt")
        self.assertFalse((self.root / "out.txt").exists())

    def test_real_legacy_xls_is_loaded_and_exported_without_changing_source(self):
        # A tiny valid raw BIFF8 workbook, so this checks xlrd rather than mocking the reader.
        def rec(code, data=b""):
            return struct.pack("<HH", code, len(data)) + data
        def bof(kind):
            return rec(0x0809, struct.pack("<HHHHII", 0x0600, kind, 0x0DBB, 0x07CC, 0, 6))
        def bounds(offset):
            return rec(0x0085, struct.pack("<IBBBB", offset, 0, 0, 4, 0) + b"Data")
        globals_data = bof(5) + rec(0x0042, struct.pack("<H", 1252)) + rec(0x0022, b"\0\0") + rec(0x00E0, b"\0" * 20)
        globals_data += bounds(len(globals_data) + len(bounds(0)) + 4) + rec(0x000A)
        sheet = bof(0x10) + rec(0x0200, struct.pack("<IIHHH", 0, 2, 0, 2, 0))
        for row, values in enumerate([["표본정보순번", "파일명"], ["29848", "photo.jpg"]]):
            for col, value in enumerate(values):
                sheet += rec(0x0204, struct.pack("<HHHHB", row, col, 0, len(value), 1) + value.encode("utf-16le"))
        data = globals_data + sheet + rec(0x000A)
        path = self.root / "legacy.xls"
        path.write_bytes(data)
        source = ExcelSource(path)
        selected = source.read_sheet("Data")
        self.assertEqual(selected.rows[0][1]["파일명"], "photo.jpg")
        profile = auto_map_profile(default_profile("specimen_image"), selected.headers)
        profile.update(schema_confirmed=True, all_rows_confirmed=True, nas_template="/NAS/images/")
        for column in profile["columns"]:
            if column["type"] == "VARCHAR2":
                column["length"] = 1000
        results = generate_sql(selected, profile)
        self.assertEqual(results[0].status, "SUCCESS", results[0].error)
        _, review = export_sql_review(source, selected, results, profile, self.root / "legacy.txt")
        workbook = load_workbook(review)
        try:
            self.assertEqual(workbook["Data"]["A2"].value, "29848")
        finally:
            workbook.close()
        self.assertEqual(path.read_bytes(), data)


if __name__ == "__main__":
    unittest.main()
