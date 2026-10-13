"""Fill the supplied NARIS templates while retaining header and cell formatting."""

from copy import copy

from openpyxl import load_workbook

from config.paths import ASSETS_DIR
from ui.widgets.column_mapping_dialog import header_key


def write_final_template(path, records, columns, kind):
    template = ASSETS_DIR / "excel" / f"{kind}_final.xlsx"
    workbook = load_workbook(template)
    try:
        sheet = workbook.active
        lookup = {header_key(target): source for target, source in columns}
        if kind == "specimen":
            lookup[header_key("이전데이터베이스번호")] = "데이터베이스번호"
            lookup[header_key("표본 공개여부")] = "표본 공개여부"
            lookup[header_key("표본활용기록")] = "표본활용현황"
        headers = [(cell.column, cell.value) for cell in sheet[2] if isinstance(cell.value, str)]
        styles = {col: copy(sheet.cell(3, col)._style) for col, _ in headers}
        for index in range(3, len(records) + 3):
            sheet.row_dimensions[index].height = sheet.row_dimensions[3].height
            for col, _ in headers:
                sheet.cell(index, col)._style = copy(styles[col])
        for col, header in headers:
            if header_key(header) == header_key("총데이터건수"):
                sheet.cell(2, col + 1, len(records))
                continue
            key = header_key(header)
            if key not in lookup:
                # The template includes optional fields absent from the input form.
                continue
            source = lookup[key]
            for index, record in enumerate(records, start=3):
                cell = sheet.cell(index, col)
                cell._style = copy(styles[col])
                cell.value = record.get(source, "") if source else ""
                if isinstance(cell.value, str):
                    # Keep dates, IDs and literal values as text (including leading '=').
                    cell.data_type = "s"
                    cell.number_format = "@"
        last_row = max(3, len(records) + 2)
        for validation in sheet.data_validations.dataValidation:
            for span in list(validation.ranges):
                if span.min_row == 3 and span.max_row < last_row:
                    span.max_row = last_row
        if sheet.max_row > last_row:
            sheet.delete_rows(last_row + 1, sheet.max_row - last_row)
        workbook.save(path)
    finally:
        workbook.close()
