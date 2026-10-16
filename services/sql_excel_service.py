"""Excel snapshots and paired SQL/review exports. No database or NAS access."""

from copy import copy
from dataclasses import dataclass
import io
import os
from pathlib import Path
import re
import shutil
import tempfile
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
import xlrd


@dataclass
class SheetData:
    name: str
    header_row: int
    headers: list[str]
    rows: list[tuple[int, dict]]


class MissingFormulaValue(str):
    """A formula whose cached result is absent; never silently insert it as NULL/text."""


def _repair_styles(data):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            value = source.read(item)
            if item.filename == "xl/styles.xml":
                counter = iter(range(100000))
                value = re.sub(rb"<((?:\w+:)?cellStyle)(?=[\s/>])(?![^>]*\bname=)",
                               lambda m: b'<%s name="Unnamed%d"' % (m.group(1), next(counter)), value)
            target.writestr(item, value)
    return output.getvalue()


class ExcelSource:
    def __init__(self, path):
        self.path = Path(path).resolve()
        if self.path.suffix.lower() not in (".xlsx", ".xls"):
            raise ValueError("xlsx 또는 xls 파일을 선택해 주세요.")
        self.data = self.path.read_bytes()
        self.legacy = self.path.suffix.lower() == ".xls"
        if not self.legacy:
            try:
                workbook = self.workbook(read_only=True)
            except TypeError as exc:
                if "_NamedCellStyle" not in str(exc):
                    raise
                self.data = _repair_styles(self.data)
                workbook = self.workbook(read_only=True)
        else:
            workbook = self.workbook()
        try:
            self.sheets = [(sheet.title, sheet.sheet_state != "visible") for sheet in workbook]
        finally:
            workbook.close()

    def workbook(self, data_only=False, read_only=False):
        if not self.legacy:
            return load_workbook(io.BytesIO(self.data), data_only=data_only, read_only=read_only)
        source = xlrd.open_workbook(file_contents=self.data, formatting_info=True)
        workbook = Workbook()
        workbook.remove(workbook.active)
        try:
            for original in source.sheets():
                sheet = workbook.create_sheet(original.name)
                sheet.sheet_state = "hidden" if original.visibility else "visible"
                for row in range(original.nrows):
                    for col in range(original.ncols):
                        cell = original.cell(row, col)
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate.xldate_as_datetime(value, source.datemode)
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif cell.ctype == xlrd.XL_CELL_EMPTY:
                            value = None
                        target = sheet.cell(row + 1, col + 1, value)
                        if isinstance(value, str):
                            target.data_type = "s"
                        if cell.xf_index < len(source.xf_list):
                            fmt = source.format_map.get(source.xf_list[cell.xf_index].format_key)
                            if fmt:
                                target.number_format = fmt.format_str
                for first_row, last_row, first_col, last_col in original.merged_cells:
                    sheet.merge_cells(start_row=first_row + 1, end_row=last_row,
                                      start_column=first_col + 1, end_column=last_col)
            return workbook
        finally:
            source.release_resources()

    def read_sheet(self, name, header_row=1):
        # SQL generation only needs values. Streaming avoids constructing all
        # styles and cell objects for large xlsx workbooks.
        workbook = self.workbook(data_only=True, read_only=not self.legacy)
        formulas = self.workbook(read_only=True) if not self.legacy else None
        try:
            sheet = workbook[name]
            if not 1 <= header_row <= sheet.max_row:
                raise ValueError("헤더 행이 시트 범위를 벗어났습니다.")
            headers, used = [], set()
            for col in range(1, sheet.max_column + 1):
                value = sheet.cell(header_row, col).value
                title = str(value).strip() if value is not None else f"열_{get_column_letter(col)}"
                unique = title if title not in used else f"{title} [{get_column_letter(col)}]"
                while unique in used:
                    unique += "_"
                used.add(unique)
                headers.append(unique)
            rows = []
            value_rows = sheet.iter_rows(min_row=header_row + 1, values_only=True)
            formula_rows = formulas[name].iter_rows(min_row=header_row + 1) if formulas is not None else None
            for number, values in enumerate(value_rows, header_row + 1):
                if formula_rows is not None:
                    formula_cells = next(formula_rows)
                    values = [MissingFormulaValue("수식 계산값 없음: 엑셀에서 재계산 후 저장 필요")
                              if value is None and cell.data_type == "f" else value
                              for value, cell in zip(values, formula_cells)]
                if any(value is not None and str(value).strip() for value in values):
                    rows.append((number, dict(zip(headers, values))))
            return SheetData(name, header_row, headers, rows)
        finally:
            workbook.close()
            if formulas is not None:
                formulas.close()

    def suggest_header_row(self, name):
        from config.sql_media_profiles import PROFILES
        def key(value):
            return re.sub(r"\s+", "", re.split(r"[(（]", str(value), maxsplit=1)[0]).casefold()
        known = {key(value) for value in ("국명", "학명", "종고유ID", "데이터베이스번호", "이전데이터베이스번호", "미디어타입", "미디어코드")}
        for profile in PROFILES.values():
            for column in profile["columns"]:
                known.update(key(value) for value in (column["name"], *column["aliases"]))
        workbook = self.workbook(data_only=True, read_only=not self.legacy)
        try:
            sheet = workbook[name]
            best_row, best_score = 1, 1
            for row in sheet.iter_rows(min_row=1, max_row=min(20, sheet.max_row)):
                score = len({key(cell.value) for cell in row if cell.value is not None} & known)
                if score > best_score:
                    best_row, best_score = row[0].row, score
            return best_row
        finally:
            workbook.close()


def sql_script(results, profile):
    statements = [result.sql for result in results if result.status == "SUCCESS"]
    lines = ["-- NARIS media INSERTs; run sequentially in one Oracle session.",
             "-- Review FK/PK constraints and NAS files separately. COMMIT is not automatic.",
             "SET DEFINE OFF", "SET SQLBLANKLINES ON", "SET AUTOCOMMIT OFF",
             "WHENEVER SQLERROR EXIT SQL.SQLCODE ROLLBACK", ""]
    if statements and profile.get("lock_table"):
        from services.sql_generation_service import identifier
        lines.append(f"LOCK TABLE {identifier(profile['table'])} IN EXCLUSIVE MODE;")
    return "\n".join(lines + statements) + "\n"


def export_sql_review(source, sheet_data, results, profile, text_path, overwrite=False):
    text_path = Path(text_path).resolve()
    if text_path.suffix.lower() not in (".txt", ".sql"):
        raise ValueError("실행용 파일 확장자는 txt 또는 sql이어야 합니다.")
    review_path = text_path.with_name(f"{text_path.stem}_검수.xlsx")
    if source.path in (text_path, review_path):
        raise ValueError("원본 엑셀을 덮어쓸 수 없습니다. 다른 파일명을 지정해 주세요.")
    if not overwrite and any(path.exists() for path in (text_path, review_path)):
        raise FileExistsError("결과 파일이 이미 존재합니다.")
    if [row for row, _ in sheet_data.rows] != [result.excel_row for result in results]:
        raise ValueError("생성 결과와 원본 행이 일치하지 않습니다. SQL을 다시 생성해 주세요.")
    if any(result.status == "SUCCESS" and len(result.sql) > 32767 for result in results):
        raise ValueError("검수용 엑셀의 SQL 셀 길이 제한(32767자)을 초과했습니다.")
    with tempfile.TemporaryDirectory(dir=text_path.parent) as folder:
        temp = Path(folder)
        sql_file, review_file = temp / "result.txt", temp / "review.xlsx"
        sql_file.write_text(sql_script(results, profile), encoding="utf-8", newline="\n")
        workbook = source.workbook()
        try:
            sheet = workbook[sheet_data.name]
            start = sheet.max_column + 1
            style = copy(sheet.cell(sheet_data.header_row, max(1, start - 1))._style)
            for offset, title in enumerate(("INSERT_SQL", "SQL_STATUS", "SQL_ERROR")):
                cell = sheet.cell(sheet_data.header_row, start + offset, title)
                cell._style = copy(style)
                cell.alignment = Alignment(vertical="center", wrap_text=False)
                sheet.column_dimensions[get_column_letter(start + offset)].width = (40, 14, 40)[offset]
            for result in results:
                values = (result.sql if result.status == "SUCCESS" else "", result.status, result.error)
                for offset, value in enumerate(values):
                    cell = sheet.cell(result.excel_row, start + offset, value)
                    cell.data_type = "s"
                    cell.alignment = Alignment(vertical="center", wrap_text=False)
            workbook.save(review_file)
        finally:
            workbook.close()
        # Prepare both files before touching either destination; restore approved existing outputs on failure.
        originals, committed = {}, []
        for index, target in enumerate((text_path, review_path)):
            if target.exists():
                backup = temp / f"backup{index}"
                shutil.copy2(target, backup)
                originals[target] = backup
        try:
            for staged, target in ((sql_file, text_path), (review_file, review_path)):
                os.replace(staged, target)
                committed.append(target)
        except OSError:
            for target in committed:
                if target in originals:
                    os.replace(originals[target], target)
                else:
                    target.unlink()
            raise
    return text_path, review_path
