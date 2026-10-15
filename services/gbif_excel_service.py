"""Export observation/specimen records into the supplied 17-column GBIF layout."""

from copy import copy
from datetime import date, datetime
from functools import lru_cache
import math

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from PySide6.QtCore import QLocale

from config.paths import ASSETS_DIR
from utils.text_utils import normalize_text


GBIF_HEADERS = (
    "occurrenceID", "basisOfRecord", "eventDate", "institutionCode",
    "scientificName", "collectionCode", "decimalLatitude", "decimalLongitude",
    "countryCode", "Identification Date", "Identified by", "stateProvince",
    "county", "locality", "verbatimLocality", "catalogNumber", "vernacularName",
)


@lru_cache(maxsize=1)
def _country_codes():
    codes = {}
    for locale in QLocale.matchingLocales(QLocale.AnyLanguage, QLocale.AnyScript, QLocale.AnyTerritory):
        code = QLocale.territoryToCode(locale.territory())
        if len(code) != 2:
            continue
        for name in (code, QLocale.territoryToString(locale.territory()), locale.nativeTerritoryName()):
            if name:
                codes[name.casefold()] = code
    for alias in ("대한민국", "한국", "Republic of Korea", "South Korea", "Korea, Republic of"):
        codes[alias.casefold()] = "KR"
    codes.update({"일본": "JP", "중국": "CN", "미국": "US", "영국": "GB", "대만": "TW"})
    return codes


def country_code(record):
    codes = _country_codes()
    # Prefer an explicitly supplied country name; never assume overseas data is KR.
    for field in ("countryCode", "국가영문명", "국가국문명"):
        value = normalize_text(record.get(field))
        if value:
            return codes.get(value.casefold(), "")
    return ""


def _date_value(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = normalize_text(value)
    if len(text) == 6 and text.isdigit():
        text = f"20{text}"
    for fmt in ("%Y-%m-%d", "%Y%m%d", "%Y.%m.%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return text


def _coordinate(value):
    text = normalize_text(value)
    try:
        number = float(text)
        return number if math.isfinite(number) else text
    except ValueError:
        return text


def gbif_values(record, kind):
    if kind not in ("observation", "specimen"):
        raise ValueError(f"Unknown GBIF record kind: {kind}")
    specimen = kind == "specimen"
    location = "채집위치" if specimen else "관찰위치"
    text = lambda field: normalize_text(record.get(field))
    return (
        text("데이터베이스번호"),
        "PreservedSpecimen" if specimen else "HumanObservation",
        _date_value(record.get("채집년월일" if specimen else "관찰일")),
        text("기관코드"), text("학명"), text("분류코드"),
        _coordinate(record.get("위도")), _coordinate(record.get("경도")),
        country_code(record), _date_value(record.get("동정년월일")),
        text("동정자 영문"), text(f"{location} 시도 영문"),
        text(f"{location} 시군구 영문"), text(f"{location} 읍면동 영문"),
        text(f"{location} 상세 영문"), text("데이터베이스번호"), text("국명"),
    )


def write_gbif_template(path, records, kind):
    workbook = load_workbook(ASSETS_DIR / "excel" / "gbif.xlsx")
    try:
        sheet = workbook.worksheets[0]
        if tuple(cell.value for cell in sheet[1]) != GBIF_HEADERS:
            raise ValueError("GBIF 양식의 컬럼 구성이 올바르지 않습니다.")
        styles = [copy(cell._style) for cell in sheet[2]]
        height = sheet.row_dimensions[2].height
        for row, record in enumerate(records, start=2):
            sheet.row_dimensions[row].height = height
            for col, value in enumerate(gbif_values(record, kind), start=1):
                cell = sheet.cell(row, col)
                cell._style = copy(styles[col - 1])
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                elif isinstance(value, date):
                    cell.number_format = "yyyy-mm-dd"
                elif col in (7, 8):
                    cell.number_format = "0.###############"
        last_row = max(2, len(records) + 1)
        if sheet.max_row > last_row:
            sheet.delete_rows(last_row + 1, sheet.max_row - last_row)
        for col, low, high in (("G", -90, 90), ("H", -180, 180)):
            validation = DataValidation(type="decimal", operator="between", formula1=low, formula2=high, allow_blank=True)
            sheet.add_data_validation(validation)
            validation.add(f"{col}2:{col}{last_row}")
        for col in ("C", "J"):
            validation = DataValidation(type="date", operator="between", formula1='DATE(1900,1,1)', formula2='DATE(9999,12,31)', allow_blank=True)
            sheet.add_data_validation(validation)
            validation.add(f"{col}2:{col}{last_row}")
        workbook.save(path)
    finally:
        workbook.close()
